"""Build Excel workbooks with embedded member-invite QR code images."""

from __future__ import annotations

import io
import secrets
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.services.guest_access_qr import qr_png_bytes_for_url

# In-memory export registry (short-lived download links).
_EXPORT_TTL_SECONDS = 60 * 60  # 1 hour
_EXPORT_LOCK = threading.Lock()
_EXPORTS: dict[str, "_ExportEntry"] = {}


@dataclass
class _ExportEntry:
    path: Path
    file_name: str
    owner_id: int
    created_at: float


@dataclass
class InviteExportRow:
    index: int
    token: str
    url: str
    expires_at: str | None
    communal_id: str | None = None


def _purge_expired_locked(now: float | None = None) -> None:
    ts = now if now is not None else time.time()
    dead = [
        key
        for key, entry in _EXPORTS.items()
        if ts - entry.created_at > _EXPORT_TTL_SECONDS
    ]
    for key in dead:
        entry = _EXPORTS.pop(key, None)
        if entry and entry.path.exists():
            try:
                entry.path.unlink()
            except OSError:
                pass


def build_member_invite_xlsx_bytes(rows: Sequence[InviteExportRow]) -> bytes:
    """Return .xlsx bytes: columns + QR image per row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Member invites"

    headers = [
        "#",
        "Communal ID",
        "Token",
        "Invite URL",
        "Expires at",
        "Use policy",
        "QR code",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 56
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 18

    # Keep image buffers alive until workbook is saved.
    image_buffers: list[io.BytesIO] = []

    for row in rows:
        excel_row = row.index + 1  # header is row 1
        expires_label = row.expires_at or "Does not expire"
        communal_label = (row.communal_id or "").strip() or "—"
        ws.append(
            [
                row.index,
                communal_label,
                row.token,
                row.url,
                expires_label,
                "Single-use",
                "",
            ]
        )
        ws.row_dimensions[excel_row].height = 90
        for col in range(1, 7):
            ws.cell(row=excel_row, column=col).alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

        png = qr_png_bytes_for_url(row.url, box_size=4, border=2)
        buf = io.BytesIO(png)
        image_buffers.append(buf)
        img = XLImage(buf)
        img.width = 96
        img.height = 96
        # Anchor in QR column (G)
        img.anchor = f"G{excel_row}"
        ws.add_image(img)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def store_invite_export(
    *,
    owner_id: int,
    content: bytes,
    file_name: str,
) -> str:
    """Persist workbook bytes and return opaque export id for download URL."""
    export_id = secrets.token_urlsafe(18)
    safe_name = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in file_name)
    if not safe_name.lower().endswith(".xlsx"):
        safe_name = f"{safe_name}.xlsx"

    export_dir = Path(temp_export_dir())
    export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / f"{export_id}.xlsx"
    path.write_bytes(content)

    with _EXPORT_LOCK:
        _purge_expired_locked()
        _EXPORTS[export_id] = _ExportEntry(
            path=path,
            file_name=safe_name,
            owner_id=int(owner_id),
            created_at=time.time(),
        )
    return export_id


def temp_export_dir() -> str:
    import tempfile

    return str(Path(tempfile.gettempdir()) / "hexzone_invite_exports")


def take_invite_export(export_id: str) -> tuple[Path, str] | None:
    """Return (path, download_file_name) if export exists and is fresh."""
    with _EXPORT_LOCK:
        _purge_expired_locked()
        entry = _EXPORTS.get(export_id)
        if not entry:
            return None
        if not entry.path.exists():
            _EXPORTS.pop(export_id, None)
            return None
        return entry.path, entry.file_name


def member_invite_join_url(token: str, *, web_base: str) -> str:
    base = (web_base or "").rstrip("/")
    return f"{base}/join?token={token}"
